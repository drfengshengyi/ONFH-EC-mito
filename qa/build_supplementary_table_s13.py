#!/usr/bin/env python3
"""Append the audited enhanced GSE123568 serum analysis as Table S13."""

from __future__ import annotations

import argparse
import hashlib
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17324D"
BLUE = "DCEAF7"
PALE = "EEF4F8"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B8C5D1")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean(value):
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def section(ws, row: int, title: str, note: str, frame: pd.DataFrame) -> int:
    ncols = max(1, len(frame.columns))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row, 1, title)
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    row += 1
    if note:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = ws.cell(row, 1, note)
        cell.font = Font(italic=True, color="465B70", size=9)
        cell.fill = PatternFill("solid", fgColor=PALE)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 32
        row += 1
    for col, name in enumerate(frame.columns, 1):
        cell = ws.cell(row, col, str(name))
        cell.font = Font(bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    row += 1
    for record in frame.itertuples(index=False, name=None):
        for col, value in enumerate(record, 1):
            cell = ws.cell(row, col, clean(value))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
        row += 1
    return row + 2


def read_csv(directory: Path, name: str) -> pd.DataFrame:
    path = directory / name
    if not path.exists():
        raise FileNotFoundError(path)
    return pd.read_csv(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-workbook", type=Path, required=True)
    parser.add_argument("--formal-dir", type=Path, required=True)
    parser.add_argument("--output-workbook", type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.source_workbook)
    if "Table S13" in workbook.sheetnames:
        workbook.remove(workbook["Table S13"])
    ws = workbook.create_sheet("Table S13")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    aggregate = read_csv(args.formal_dir, "aggregate_performance.csv")
    aggregate.insert(1, "n_samples", 40)
    aggregate.insert(2, "n_sonfh", 30)
    aggregate.insert(3, "n_control", 10)
    aggregate.insert(4, "n_genes", aggregate["model"].map({"candidate": 77, "ma": 4}))
    aggregate["no_skill_ap"] = 0.75
    aggregate["no_skill_brier"] = 0.1875

    repeats = read_csv(args.formal_dir, "repeat_performance.csv")
    permutation = read_csv(args.formal_dir, "label_permutation_summary.csv").iloc[0]
    permutation_null = read_csv(args.formal_dir, "label_permutation_null.csv")
    random_summary = read_csv(args.formal_dir, "matched_random_space_summary.csv").iloc[0]
    random_null = read_csv(args.formal_dir, "matched_random_space_auc_null.csv")
    nulls = pd.DataFrame([
        {
            "test": "full-pipeline label permutation",
            "observed_auc": permutation.observed_auc,
            "n_null": int(permutation.n_valid),
            "n_null_ge_observed": int(permutation.n_ge_observed),
            "empirical_p": permutation.empirical_p,
            "null_auc_q025": permutation_null.auc.quantile(0.025),
            "null_auc_median": permutation_null.auc.median(),
            "null_auc_q975": permutation_null.auc.quantile(0.975),
            "scope": "class-size-preserving permutations; full candidate pipeline rerun",
        },
        {
            "test": "expression/variance-matched random 77-gene spaces",
            "observed_auc": random_summary.observed_candidate_auc,
            "n_null": int(random_summary.n_matched_spaces),
            "n_null_ge_observed": int(random_summary.n_ge_observed),
            "empirical_p": random_summary.empirical_p_specificity,
            "null_auc_q025": random_null.auc.quantile(0.025),
            "null_auc_median": random_null.auc.median(),
            "null_auc_q975": random_null.auc.quantile(0.975),
            "scope": random_summary.matching_scope,
        },
    ])

    delong = read_csv(args.formal_dir, "delong_comparison.csv")
    fixed = read_csv(args.formal_dir, "fixed_oof_paired_bootstrap_summary.csv")
    refit = read_csv(args.formal_dir, "refit_oob_bootstrap_summary.csv")
    comparisons = pd.concat([
        pd.DataFrame([{
            "analysis": "paired DeLong on fixed aggregate OOF predictions",
            "metric": "delta_auc",
            "estimate": delong.loc[0, "delta_auc"],
            "q025": None,
            "q975": None,
            "p_value": delong.loc[0, "paired_delong_p"],
            "n_valid": 1,
            "scope": delong.loc[0, "scope"],
        }]),
        fixed.rename(columns={"median": "estimate", "two_sided_zero_p": "p_value"}).assign(
            analysis="fixed-OOF paired stratified bootstrap", n_valid=20000
        )[["analysis", "metric", "estimate", "q025", "q975", "p_value", "n_valid", "scope"]],
        refit.rename(columns={"median": "estimate", "interval_name": "scope"}).assign(
            analysis="refit-aware stratified OOB bootstrap", p_value=None
        )[["analysis", "metric", "estimate", "q025", "q975", "p_value", "n_valid", "scope"]],
    ], ignore_index=True)

    stability = read_csv(args.formal_dir, "feature_stability_summary.csv")
    stability_columns = [
        "gene", "n_outer_fits", "n_nonzero", "nonzero_frequency", "n_positive", "n_negative",
        "dominant_direction", "direction_consistency_given_selected", "coefficient_median_all",
        "coefficient_q1_all", "coefficient_q3_all", "coefficient_median_nonzero",
        "coefficient_q1_nonzero", "coefficient_q3_nonzero", "stable_frequency_ge_0_60",
        "stable_frequency_and_direction",
    ]
    stability = stability[stability_columns].sort_values(
        ["nonzero_frequency", "gene"], ascending=[False, True]
    )
    selected_counts = read_csv(args.formal_dir, "features_per_outer_fit.csv")
    calibration = read_csv(args.formal_dir, "calibration_summary.csv")
    calibration_bins = read_csv(args.formal_dir, "calibration_bins.csv")
    legacy = read_csv(args.formal_dir, "legacy_vs_enhanced_comparison.csv")
    invalid = read_csv(args.formal_dir, "refit_oob_bootstrap_invalid.csv")

    design = pd.DataFrame([
        ("sample set", "GSE123568; 30 SONFH and 10 steroid-exposed non-SONFH controls"),
        ("candidate space", "77 prespecified measurable genes"),
        ("comparator", "Ma et al.; BID, FTH1, LACTB and PDK3"),
        ("outer CV", "stratified 5 folds x 5 repeats; random_state 0--4"),
        ("inner CV", "stratified 4 folds; AUC-selected C"),
        ("regularization grid", "16 log-spaced C values from 1e-3 to 1e2"),
        ("probe selection", "highest mean in each outer training fold; exact ties by lexicographic probe ID"),
        ("refit-aware OOB", f"1000 valid runs; {len(invalid)} invalid attempts before target"),
        ("base seed", "20260820"),
    ], columns=["item", "frozen specification"])

    source_names = [
        "aggregate_performance.csv", "repeat_performance.csv", "label_permutation_summary.csv",
        "label_permutation_null.csv", "matched_random_space_summary.csv",
        "matched_random_space_auc_null.csv", "matched_random_space_qc.csv",
        "delong_comparison.csv", "fixed_oof_paired_bootstrap_summary.csv",
        "refit_oob_bootstrap_summary.csv", "refit_oob_bootstrap_invalid.csv",
        "feature_stability_summary.csv", "features_per_outer_fit.csv", "calibration_summary.csv",
        "calibration_bins.csv", "legacy_vs_enhanced_comparison.csv",
    ]
    audit_rows = []
    for name in source_names:
        path = args.formal_dir / name
        frame = pd.read_csv(path)
        audit_rows.append({
            "source_file": name,
            "rows": len(frame),
            "columns": len(frame.columns),
            "sha256": sha256(path),
        })
    audit = pd.DataFrame(audit_rows)

    row = 1
    row = section(ws, row, "Supplementary Table S13. Enhanced GSE123568 serum-analysis audit", "All metrics are exploratory internal cross-validated estimates from the same 40 samples. The table replaces the earlier serum summary but does not constitute external validation, a locked signature or a clinical threshold.", design)
    row = section(ws, row, "Table S13a. Aggregate performance", "Each sample's five out-of-fold probabilities were averaged before aggregate metrics were calculated.", aggregate)
    row = section(ws, row, "Table S13b. Repeat-specific performance", "Repeat-level values are paired descriptive outputs; they were not treated as independent replicates.", repeats)
    row = section(ws, row, "Table S13c. Null-model calibration", "Empirical p=(1 + number of null AUCs at least as large as observed)/(1 + number of null runs).", nulls)
    row = section(ws, row, "Table S13d. Candidate-versus-Ma paired comparisons", "Fixed-OOF intervals exclude refitting variability; OOB intervals repeat the modeling pipeline but remain internal to the same dataset.", comparisons)
    row = section(ws, row, "Table S13e. Candidate feature-selection stability", "Coefficient direction is conditional on the multivariable training-fold model and is not a marginal expression effect.", stability)
    row = section(ws, row, "Table S13f. Number of selected genes per outer fit", "Nonzero coefficients from the L1 candidate model.", selected_counts)
    row = section(ws, row, "Table S13g. Exploratory calibration summaries", "Calibration intercepts, slopes and four-bin summaries are unstable at n=40 and are not evidence of clinical calibration.", calibration)
    row = section(ws, row, "Table S13h. Exploratory calibration bins", "Each probability-ranked bin contains 10 samples.", calibration_bins)
    row = section(ws, row, "Table S13i. Legacy-to-enhanced analysis audit", "The enhanced workflow moves probe selection into outer training folds and adds matched spaces and refit-aware OOB uncertainty.", legacy)
    section(ws, row, "Table S13j. Source-file integrity manifest", "SHA-256 values refer to the frozen formal output directory serum_enhanced_20260910/02_formal.", audit)

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_length = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_length + 2, 11), 34)
    ws.column_dimensions["A"].width = 46

    # Remove a now-stale note in the copied Table S12.
    if "Table S12" in workbook.sheetnames:
        workbook["Table S12"]["A2"] = (
            "These analyses do not replace the frozen primary workflow. They quantify dependence on ComBat use, marker-panel margin, and the 26-EC HOA1 sampling unit. A separate enhanced serum analysis is reported in Table S13."
        )

    args.output_workbook.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_workbook)
    print(f"WROTE {args.output_workbook} with sheets={len(workbook.sheetnames)} rows_Table_S13={ws.max_row}")


if __name__ == "__main__":
    main()
