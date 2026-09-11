"""Append the Scientific Reports preprocessing audits as Supplementary Table S12."""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SECTIONS = [
    (
        "Table S12a. Atlas no-ComBat summary",
        "atlas_no_combat_summary.csv",
        "The archived uncorrected 30-component PCA was used to rebuild the neighbour graph, Leiden clusters and identical marker-panel assignments.",
    ),
    (
        "Table S12b. Sampling-unit EC fractions with and without ComBat",
        "atlas_no_combat_ec_composition.csv",
        "The no-ComBat solution is an alternative preprocessing result, not a batch-free reference standard.",
    ),
    (
        "Table S12c. EC-state no-ComBat summary",
        "no_combat_summary.csv",
        "The fixed 13,426-cell primary EC set was re-embedded, clustered and panel-labelled without ComBat.",
    ),
    (
        "Table S12d. Type H participant-level statistics after EC-state no-ComBat reanalysis",
        "no_combat_typeH_exact_statistics.csv",
        "Exact tests use independent Liao participants; effects are percentage-point differences.",
    ),
    (
        "Table S12e. Annotation-margin threshold scan summary",
        "annotation_margin_threshold_summary.csv",
        "Clusters with top-two panel-score margins below each threshold were labelled ambiguous.",
    ),
    (
        "Table S12f. Annotation-margin threshold exact statistics",
        "annotation_margin_threshold_exact_statistics.csv",
        "No threshold was selected from these results; the complete prespecified scan is shown.",
    ),
    (
        "Table S12g. Sampling-unit depth and pseudobulk inclusion audit",
        "pseudobulk_sampling_unit_depth_audit.csv",
        "stored_ec_nCount_sum is the sum of the per-cell nCount field archived in the frozen EC object; it is provided as a depth audit and is not used as an independent replicate.",
    ),
    (
        "Table S12h. Complete omit-HOA1 refit for the four prioritized Hallmark pathways",
        "omit_hoa1_hallmark_refit.csv",
        "This is the existing full differential-expression and enrichment refit with the 26-EC HOA1 pseudobulk omitted; the four rows are the prespecified pathways shown in Figure 4B.",
    ),
]


def clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--results", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    if "Table S12" in workbook.sheetnames:
        del workbook["Table S12"]
    sheet = workbook.create_sheet("Table S12")
    sheet.sheet_view.showGridLines = False

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="EAF2F8")
    dark_font = Font(color="FFFFFF", bold=True, size=12)
    section_font = Font(color="17365D", bold=True, size=11)
    header_font = Font(color="17365D", bold=True)

    max_columns = 16
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
    sheet.cell(1, 1, "Supplementary Table S12. Preprocessing, annotation-threshold, and low-depth pseudobulk sensitivity audits")
    sheet.cell(1, 1).fill = title_fill
    sheet.cell(1, 1).font = dark_font
    sheet.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_columns)
    sheet.cell(2, 1, "These analyses do not replace the frozen primary workflow. They quantify dependence on ComBat use, marker-panel margin, and the 26-EC HOA1 sampling unit. A separate enhanced serum analysis is reported in Table S13.")
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 34

    row = 4
    for section_index, (title, filename, note) in enumerate(SECTIONS, start=1):
        path = args.results / filename
        if not path.exists():
            raise FileNotFoundError(path)
        frame = pd.read_csv(path)
        if filename == "omit_hoa1_hallmark_refit.csv":
            keep = ["fit_id", "pathway", "NES", "pval", "padj", "size", "leading_edge_count", "leading_edge_genes"]
            frame = frame[[column for column in keep if column in frame.columns]]

        last_column = max(1, len(frame.columns))
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
        cell = sheet.cell(row, 1, title)
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        row += 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
        sheet.cell(row, 1, note).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 30
        row += 1

        header_row = row
        for column_index, column in enumerate(frame.columns, start=1):
            cell = sheet.cell(row, column_index, column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        row += 1
        for values in frame.itertuples(index=False, name=None):
            for column_index, value in enumerate(values, start=1):
                cell = sheet.cell(row, column_index, clean_value(value))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000E+00" if abs(cell.value) < 0.001 and cell.value != 0 else "0.0000"
            row += 1
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(last_column)}{row - 1}"
        row += 2

    sheet.freeze_panes = "A4"
    for column in range(1, max_columns + 1):
        letter = get_column_letter(column)
        values = [sheet.cell(r, column).value for r in range(1, sheet.max_row + 1)]
        width = min(42, max(11, max((len(str(value)) for value in values if value is not None), default=0) + 2))
        sheet.column_dimensions[letter].width = width

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    check = load_workbook(args.output, read_only=True, data_only=False)
    if "Table S12" not in check.sheetnames:
        raise RuntimeError("Table S12 missing after workbook save")
    print(f"Saved {args.output} with sheets: {', '.join(check.sheetnames)}")


if __name__ == "__main__":
    main()
